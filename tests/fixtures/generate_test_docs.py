"""Generates test .docx fixtures (spec §31):
    correct.docx, wrong_font.docx, wrong_margins.docx, wrong_spacing.docx,
    missing_sections.docx, mixed_formatting.docx

Run directly: `python -m tests.fixtures.generate_test_docs`
Also importable — pytest fixtures in tests/conftest.py call build_* functions
directly so the .docx files are (re)generated fresh for every test run
rather than committed as binary blobs to git.
"""

from __future__ import annotations

from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Mm, Pt

FIXTURES_DIR = Path(__file__).parent

_BODY_PARAGRAPHS = {
    "Введение": (
        "Актуальность настоящего исследования обусловлена растущим интересом "
        "к автоматизации проверки учебных работ. Проблема заключается в том, "
        "что ручная проверка требует значительных временных затрат. "
        "Цель работы — разработать систему автоматической проверки курсовых "
        "работ. Задачи: изучить предметную область, спроектировать архитектуру, "
        "реализовать прототип. Объектом исследования являются курсовые работы "
        "студентов. Предметом исследования являются методы автоматической "
        "проверки текста. В работе использованы методы анализа документов и "
        "обработки естественного языка."
    ),
    "Теоретическая часть": (
        "В данном разделе рассматриваются теоретические основы автоматической "
        "проверки документов. Существующие подходы делятся на два класса: "
        "детерминированные и основанные на машинном обучении. Первые "
        "обеспечивают стопроцентную точность для измеримых параметров, "
        "вторые — гибкость при анализе смысла текста."
    ),
    "Практическая часть": (
        "В практической части описана реализация системы проверки. "
        "Система состоит из модуля разбора документа, модуля правил и "
        "модуля искусственного интеллекта. Проведено тестирование на "
        "наборе документов различного качества оформления."
    ),
    "Заключение": (
        "В ходе выполнения работы были решены все поставленные задачи. "
        "Разработана архитектура системы автоматической проверки курсовых "
        "работ. Реализован прототип, сочетающий детерминированные проверки "
        "и анализ текста с помощью искусственного интеллекта. Результаты "
        "тестирования подтвердили работоспособность предложенного подхода."
    ),
}

_REFERENCES = [
    f"{i}. Автор {i}. Название работы {i}. Издательство, 20{10 + i % 9}. — {100 + i} с."
    for i in range(1, 16)
]


def _set_page_and_margins(document, top=20, bottom=20, left=30, right=10):
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(top)
    section.bottom_margin = Mm(bottom)
    section.left_margin = Mm(left)
    section.right_margin = Mm(right)


def _add_heading(document, text, level=1):
    heading = document.add_heading(text, level=level)
    return heading


def _add_body_paragraph(
    document, text, font_name="Times New Roman", size_pt=14, indent_cm=1.25, spacing=1.5
):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = p.paragraph_format
    fmt.first_line_indent = Cm(indent_cm)
    fmt.line_spacing = spacing
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    return p


def build_correct() -> Path:
    document = docx.Document()
    _set_page_and_margins(document)
    for section, text in _BODY_PARAGRAPHS.items():
        _add_heading(document, section)
        _add_body_paragraph(document, text)
    _add_heading(document, "Список литературы")
    for entry in _REFERENCES:
        _add_body_paragraph(document, entry, indent_cm=0)

    # in-text citations
    _add_body_paragraph(document, "Данный подход подтверждается рядом исследований [1, 2, 3].")

    path = FIXTURES_DIR / "correct.docx"
    document.save(str(path))
    return path


def build_wrong_font() -> Path:
    document = docx.Document()
    _set_page_and_margins(document)
    for section, text in _BODY_PARAGRAPHS.items():
        _add_heading(document, section)
        _add_body_paragraph(document, text, font_name="Arial", size_pt=12)
    _add_heading(document, "Список литературы")
    for entry in _REFERENCES:
        _add_body_paragraph(document, entry, font_name="Arial", size_pt=12, indent_cm=0)
    path = FIXTURES_DIR / "wrong_font.docx"
    document.save(str(path))
    return path


def build_wrong_margins() -> Path:
    document = docx.Document()
    _set_page_and_margins(document, top=10, bottom=10, left=15, right=15)
    for section, text in _BODY_PARAGRAPHS.items():
        _add_heading(document, section)
        _add_body_paragraph(document, text)
    _add_heading(document, "Список литературы")
    for entry in _REFERENCES:
        _add_body_paragraph(document, entry, indent_cm=0)
    path = FIXTURES_DIR / "wrong_margins.docx"
    document.save(str(path))
    return path


def build_wrong_spacing() -> Path:
    document = docx.Document()
    _set_page_and_margins(document)
    for section, text in _BODY_PARAGRAPHS.items():
        _add_heading(document, section)
        _add_body_paragraph(document, text, spacing=1.0)
    _add_heading(document, "Список литературы")
    for entry in _REFERENCES:
        _add_body_paragraph(document, entry, indent_cm=0, spacing=1.0)
    path = FIXTURES_DIR / "wrong_spacing.docx"
    document.save(str(path))
    return path


def build_missing_sections() -> Path:
    document = docx.Document()
    _set_page_and_margins(document)
    # Deliberately omit "Практическая часть" and "Список литературы".
    for section in ("Введение", "Теоретическая часть", "Заключение"):
        _add_heading(document, section)
        _add_body_paragraph(document, _BODY_PARAGRAPHS[section])
    path = FIXTURES_DIR / "missing_sections.docx"
    document.save(str(path))
    return path


def build_mixed_formatting() -> Path:
    document = docx.Document()
    _set_page_and_margins(document)
    fonts = [("Times New Roman", 14), ("Arial", 12), ("Calibri", 11)]
    for i, (section, text) in enumerate(_BODY_PARAGRAPHS.items()):
        _add_heading(document, section)
        font_name, size = fonts[i % len(fonts)]
        _add_body_paragraph(document, text, font_name=font_name, size_pt=size)
    _add_heading(document, "Список литературы")
    for entry in _REFERENCES:
        _add_body_paragraph(document, entry, indent_cm=0)
    path = FIXTURES_DIR / "mixed_formatting.docx"
    document.save(str(path))
    return path


def build_all() -> list[Path]:
    return [
        build_correct(),
        build_wrong_font(),
        build_wrong_margins(),
        build_wrong_spacing(),
        build_missing_sections(),
        build_mixed_formatting(),
    ]


if __name__ == "__main__":
    for p in build_all():
        print(f"generated {p}")
