import json
import zipfile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt
from pydantic import ValidationError

from app.ai.analyzer import AIAnalyzer
from app.ai.provider import LLMProvider, LLMResponse, MockProvider
from app.ai.schemas import AIErrorList, IntroductionAnalysis
from app.analysis.aggregator import aggregate
from app.bot.messages import split_message
from app.common.utils import sanitize_display_name
from app.config.settings import Settings
from app.document.parser import parse_docx
from app.document.structure import analyze_structure, split_into_logical_sections
from app.rules.engine import RuleEngine
from app.rules.validators.font import validate_font
from app.rules.validators.paragraphs import validate_paragraphs
from app.rules.validators.references import validate_references
from app.rules.validators.spacing import validate_spacing
from app.security.validation import validate_docx_container


def saved(document, tmp_path):
    path = tmp_path / "document.docx"
    document.save(path)
    return parse_docx(path)


def test_inherited_formatting_and_explicit_zero(tmp_path, sample_preset):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(14)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = 1.5
    p = doc.add_paragraph("Inherited body")
    parsed = saved(doc, tmp_path)
    assert parsed.dominant_font == ("Times New Roman", 14)
    assert all(f.severity.value == "pass" for f in validate_font(parsed, sample_preset))
    assert all(f.severity.value == "pass" for f in validate_paragraphs(parsed, sample_preset))
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    parsed = saved(doc, tmp_path)
    assert parsed.paragraphs[0].alignment == "left"
    assert any(f.severity.value == "error" for f in validate_paragraphs(parsed, sample_preset))


@pytest.mark.parametrize(
    "spacing,expected", [(1.0, "error"), (1.5, "pass"), (2.0, "error"), (Pt(30), "error")]
)
def test_spacing_modes(tmp_path, sample_preset, spacing, expected):
    doc = Document()
    doc.add_paragraph("Body").paragraph_format.line_spacing = spacing
    assert validate_spacing(saved(doc, tmp_path), sample_preset)[0].severity.value == expected


def test_custom_heading_inherited_outline(tmp_path):
    doc = Document()
    base = doc.styles.add_style("Custom base", WD_STYLE_TYPE.PARAGRAPH)
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "0")
    base.element.get_or_add_pPr().append(outline)
    child = doc.styles.add_style("Custom child", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = base
    doc.add_paragraph("Introduction", child)
    assert saved(doc, tmp_path).paragraphs[0].is_heading


def test_unheaded_and_table_text_reaches_sections(tmp_path):
    doc = Document()
    doc.add_table(1, 1).cell(0, 0).text = "Text inside a table"
    parsed = saved(doc, tmp_path)
    sections = split_into_logical_sections(parsed, analyze_structure(parsed))
    assert "Text inside a table" in sections["body"]


def test_all_sections_are_checked(tmp_path, sample_preset):
    doc = Document()
    doc.add_paragraph("Body")
    doc.sections[0].left_margin = Mm(30)
    doc.add_section().left_margin = Mm(50)
    findings = RuleEngine().run(saved(doc, tmp_path), sample_preset).findings
    assert any(f.rule_id == "margins.left_margin" and f.actual == "50 мм" for f in findings)


def test_numbered_duplicate_sources(tmp_path, sample_preset):
    doc = Document()
    doc.add_heading("References", 1)
    doc.add_paragraph("1. Author. Identical book. 2020.")
    doc.add_paragraph("2. Author. Identical book. 2020.")
    parsed = saved(doc, tmp_path)
    structure = analyze_structure(parsed)
    findings = validate_references(
        split_into_logical_sections(parsed, structure)["references"], structure, sample_preset
    )
    assert any(f.rule_id == "references.duplicates" for f in findings)


def test_archive_expansion_and_duplicate_names(tmp_path):
    path = tmp_path / "compressed.docx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "x" * 1_000_000)
    assert not validate_docx_container(path).ok


def test_long_filename_preserves_extension():
    assert sanitize_display_name("a" * 200 + ".docx").endswith(".docx")
    assert len(sanitize_display_name("a" * 200 + ".docx")) == 120


def test_message_chunks_handle_unicode():
    text = "😀" * 5000
    parts = split_message(text)
    assert "".join(parts) == text
    assert all(len(p.encode("utf-16-le")) // 2 <= 3500 for p in parts)


def test_empty_ai_objects_are_invalid():
    with pytest.raises(ValidationError):
        AIErrorList.model_validate({})
    with pytest.raises(ValidationError):
        IntroductionAnalysis.model_validate({"errors": []})


@pytest.mark.asyncio
async def test_mock_never_claims_analysis():
    analyzer = AIAnalyzer(MockProvider(), Settings(_env_file=None))
    assert await analyzer.analyze({"introduction": "text"}) == ([], False)


class RecordingProvider(LLMProvider):
    name = "recording"

    def __init__(self):
        self.calls = 0

    async def complete(self, **kwargs):
        self.calls += 1
        return LLMResponse(json.dumps({"errors": []}), "test")


@pytest.mark.asyncio
async def test_budget_includes_each_request_and_output():
    settings = Settings(_env_file=None, MAX_AI_TOKENS_PER_CHECK=4000)
    provider = RecordingProvider()
    analyzer = AIAnalyzer(provider, settings)
    findings, ok = await analyzer.analyze({"body": "A paragraph"})
    assert not ok
    assert provider.calls < 3
    assert analyzer.tokens_used <= 4000


def test_unavailable_categories_do_not_get_points(sample_preset):
    result = aggregate([], [], sample_preset, {}, False)
    assert result.max_score == 50
    assert result.score == 50
    assert all(
        c.earned_points == 0 and not c.evaluated
        for c in result.category_scores
        if c.category.value in {"language", "style", "content"}
    )


def test_preset_rules_are_executed(tmp_path, sample_preset):
    doc = Document()
    doc.add_paragraph("Cover")
    doc.add_heading("Introduction", 1)
    doc.add_paragraph("Short text")
    sample_preset.headings.require_page_break_before_top_level = True
    rules = {f.rule_id for f in RuleEngine().run(saved(doc, tmp_path), sample_preset).findings}
    assert "headings.page_break" in rules
    assert "structure.min_words.introduction" in rules


def test_settings_reject_unknown_provider_and_negative_limit():
    with pytest.raises(ValidationError):
        Settings(_env_file=None, LLM_PROVIDER="typo")
    with pytest.raises(ValidationError):
        Settings(_env_file=None, MAX_PAGES=-1)


def test_pdf_fallback_is_stable(monkeypatch):
    from app.reports import fonts

    monkeypatch.setattr(fonts, "_registered", False)
    monkeypatch.setattr(fonts, "_registered_family", "Helvetica")
    monkeypatch.setattr(fonts, "_CANDIDATES", [])
    assert fonts.ensure_unicode_font_registered() == "Helvetica"
    assert fonts.ensure_unicode_font_registered() == "Helvetica"
